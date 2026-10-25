"""M5 验收：材料上传、文本抽取、归属鉴权、单章重跑。

跑法（用项目虚拟环境）：
    .venv/Scripts/python.exe tests/m5_materials.py

这个用例会用真实模型跑一次短生成（约 40 秒）与一次单章重跑（约 15 秒），
因为「重跑后是否真的重新落盘」只能由真实链路证明。
"""
from __future__ import annotations

import json
import os
import sqlite3
import sys
import time
import urllib.error
import urllib.request
import uuid
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402

BASE = os.environ.get("DA_BASE", "http://127.0.0.1:8020")
DB_PATH = ROOT / "data" / "app.db"
TMP = ROOT / "data" / "tmp" / "m5"
PASS, FAIL = [], []


def check(name: str, ok: bool, detail: str = "") -> None:
    (PASS if ok else FAIL).append(name)
    print(f"{'  OK ' if ok else '  !! '} {name}" + (f"  [{detail}]" if detail else ""))


def read_env() -> dict:
    env = {}
    for line in (ROOT / ".env").read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            env[key.strip()] = value.strip().strip('"').strip("'")
    return env


# ---------------- 构造样本文件 ----------------

def make_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("受案范围的三次收缩", level=1)
    doc.add_paragraph(
        "行政协议进入受案范围之后，司法解释对可诉类型作了三次收缩。"
        "每一次收缩都伴随着一个新的判断标准，而这些标准彼此之间并不完全一致。"
    )
    doc.add_paragraph(
        "这一不一致直接导致同类案件在不同法院得到不同处理，"
        "也给当事人选择救济途径带来了不确定性。"
    )
    doc.save(str(path))


def make_pdf(path: Path, text: str | None) -> None:
    """手搓一个最小 PDF：text 为 None 时只画一个矩形，形成「无文字层」的扫描件效果。"""
    if text:
        escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET"
    else:
        content = "0.8 0.8 0.8 rg 72 600 400 120 re f"
    stream = content.encode("latin-1", "replace")

    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 595 842]"
        b"/Resources<</Font<</F1 5 0 R>>>>/Contents 4 0 R>>",
        b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (f"trailer\n<</Size {len(objects) + 1}/Root 1 0 R>>\nstartxref\n"
            f"{xref_at}\n%%EOF\n").encode()
    path.write_bytes(bytes(out))


# ---------------- HTTP ----------------

class Client:
    def __init__(self):
        import http.cookiejar

        self.jar = http.cookiejar.CookieJar()
        self.opener = urllib.request.build_opener(
            urllib.request.ProxyHandler({}),  # 关掉本机 HTTP_PROXY
            urllib.request.HTTPCookieProcessor(self.jar),
        )
        self.xff = f"10.41.0.{uuid.uuid4().int % 200 + 1}"

    def call(self, method: str, path: str, body=None):
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(BASE + path, data=data, method=method)
        req.add_header("Content-Type", "application/json")
        req.add_header("X-Forwarded-For", self.xff)
        try:
            with self.opener.open(req, timeout=300) as resp:
                raw = resp.read()
                ctype = resp.headers.get("content-type", "")
                if "json" in ctype:
                    return resp.status, json.loads(raw.decode("utf-8", "replace"))
                return resp.status, raw
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode("utf-8", "replace")
            try:
                return exc.code, json.loads(raw)
            except json.JSONDecodeError:
                return exc.code, raw

    def upload(self, task_id: str, paths: list[Path]):
        boundary = "----da" + uuid.uuid4().hex
        body = bytearray()
        for path in paths:
            body += f"--{boundary}\r\n".encode()
            body += (f'Content-Disposition: form-data; name="files"; filename="{path.name}"\r\n'
                     f"Content-Type: application/octet-stream\r\n\r\n").encode()
            body += path.read_bytes() + b"\r\n"
        body += f"--{boundary}--\r\n".encode()
        req = urllib.request.Request(f"{BASE}/api/tasks/{task_id}/files", data=bytes(body), method="POST")
        req.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
        req.add_header("X-Forwarded-For", self.xff)
        try:
            with self.opener.open(req, timeout=120) as resp:
                return resp.status, json.loads(resp.read().decode("utf-8", "replace"))
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode("utf-8", "replace")
            try:
                return exc.code, json.loads(raw)
            except json.JSONDecodeError:
                return exc.code, raw

    def stream(self, path: str, body=None):
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(BASE + path, data=data, method="POST")
        req.add_header("Content-Type", "application/json")
        req.add_header("Accept", "text/event-stream")
        req.add_header("X-Forwarded-For", self.xff)
        try:
            with self.opener.open(req, timeout=600) as resp:
                return resp.status, parse_sse(resp.read().decode("utf-8", "replace"))
        except urllib.error.HTTPError as exc:
            return exc.code, []


def parse_sse(raw: str) -> list[tuple[str, dict]]:
    events = []
    for block in raw.split("\n\n"):
        if not block.strip() or block.strip().startswith(":"):
            continue
        name, payload = "message", {}
        for line in block.splitlines():
            if line.startswith("event:"):
                name = line[6:].strip()
            elif line.startswith("data:"):
                try:
                    payload = json.loads(line[5:].strip())
                except json.JSONDecodeError:
                    payload = {}
        events.append((name, payload))
    return events


def new_guest(admin: Client, nickname: str) -> tuple[Client, dict]:
    status, body = admin.call("POST", "/api/admin/invites",
                              {"count": 1, "max_uses": 1, "days": 7, "note": f"m5-{nickname}"})
    code = (body.get("codes") or [""])[0]
    guest = Client()
    guest.call("POST", "/api/join", {"code": code, "nickname": nickname})
    status, me = guest.call("GET", "/api/me")
    return guest, me


def new_task(guest: Client, goal: str, mode: str = "writing") -> str:
    status, body = guest.call("POST", "/api/tasks", {"mode": mode, "goal": goal})
    return ((body or {}).get("task") or {}).get("id") or ""


def main() -> int:
    TMP.mkdir(parents=True, exist_ok=True)
    admin = Client()
    status, _ = admin.call("POST", "/api/admin/login",
                           {"password": read_env().get("ADMIN_PASSWORD", "")})
    if status != 200:
        print("管理员登录失败")
        return 2

    guest, me = new_guest(admin, "m5甲")
    task_id = new_task(guest, "读这几份材料，写一份关于受案范围收缩的短文")
    check("建任务", bool(task_id), task_id)

    print("\n[1] 四类材料抽取")
    md = TMP / "notes.md"
    md.write_text(
        "# 受案范围笔记\n\n司法解释把可诉的行政协议类型逐一列出，列举之外并非当然不可诉。\n\n"
        "关键在于协议是否具备行政法上的权利义务内容，这一点比类型清单更难操作。\n",
        encoding="utf-8",
    )
    docx = TMP / "chapter.docx"
    make_docx(docx)
    pdf_text = TMP / "with_text.pdf"
    make_pdf(pdf_text, "Administrative agreement litigation concerns the scope of "
                       "acceptable cases and the standard of legality review in Chinese courts.")
    pdf_scan = TMP / "scan_only.pdf"
    make_pdf(pdf_scan, None)

    status, body = guest.upload(task_id, [md, docx, pdf_text, pdf_scan])
    check("上传 4 份回 200", status == 200, f"got {status} {str(body)[:120]}")
    saved = {item["name"]: item for item in (body.get("files") or [])}
    check("四份都登记了", len(saved) == 4, str(list(saved)))

    if "notes.md" in saved:
        meta = saved["notes.md"]["meta"]
        check("Markdown 抽到文字", meta["has_text"] and meta["chars"] > 40,
              f"chars={meta['chars']}")
    if "chapter.docx" in saved:
        meta = saved["chapter.docx"]["meta"]
        check("DOCX 抽到文字", meta["has_text"] and meta["chars"] > 40, f"chars={meta['chars']}")
    if "with_text.pdf" in saved:
        meta = saved["with_text.pdf"]["meta"]
        check("PDF 有文字层", meta["has_text"], f"pages={meta.get('page_count')} chars={meta['chars']}")
    if "scan_only.pdf" in saved:
        meta = saved["scan_only.pdf"]["meta"]
        check("无文字层 PDF 被判为扫描件", meta["has_text"] is False and "扫描件" in (meta.get("note") or ""),
              str(meta.get("note"))[:40])

    sidecar = ROOT / "data" / "extracted" / f"{saved.get('notes.md', {}).get('id', 'x')}.txt"
    check("抽取文本落了旁文件", sidecar.exists() and sidecar.stat().st_size > 20,
          f"{sidecar.name} {sidecar.stat().st_size if sidecar.exists() else 0} 字节")

    print("\n[2] 类型与内容校验")
    fake = TMP / "fake.docx"
    fake.write_text("这其实是纯文本，不是 docx", encoding="utf-8")
    status, body = guest.upload(task_id, [fake])
    check("扩展名伪装被拒 400", status == 400 and "真实格式" in json.dumps(body, ensure_ascii=False),
          f"got {status} {str(body)[:80]}")

    bad = TMP / "tool.exe"
    bad.write_bytes(b"MZ\x90\x00")
    status, body = guest.upload(task_id, [bad])
    check("不允许的扩展名被拒 400", status == 400 and "不支持" in json.dumps(body, ensure_ascii=False),
          f"got {status} {str(body)[:80]}")

    print("\n[3] 归属鉴权")
    other, other_me = new_guest(admin, "m5乙")
    other_task = new_task(other, "别人的任务")
    status, body = other.call("GET", f"/api/files/{saved.get('notes.md', {}).get('id', 'x')}")
    check("他人下载材料被拒", status in (403, 404), f"got {status}")
    status, body = other.upload(task_id, [md])
    check("他人向我的任务上传被拒 403", status == 403, f"got {status}")
    status, body = other.stream(f"/api/tasks/{task_id}/generate", {"word_count": 300})
    check("他人生成我的任务被拒 403", status == 403, f"got {status}")
    status, body = other.call("GET", f"/api/tasks/{task_id}")
    check("他人读我的任务被拒 403", status == 403, f"got {status}")

    print("\n[4] 材料确实进了提示词")
    status, events = guest.stream(f"/api/tasks/{task_id}/chat", {"text": '#fake {"chunks": 1, "delay_ms": 10}'})
    meta = next((p for n, p in events if n == "meta"), {})
    check("meta 报告材料份数", (meta.get("materials") or 0) >= 3, f"materials={meta.get('materials')}")

    print("\n[5] 文件列表按类型分")
    status, body = guest.call("GET", "/api/files?kind=source")
    kinds = {item["kind"] for item in (body.get("files") or [])}
    check("source 列表只含材料", kinds == {"source"}, str(kinds))

    print("\n[6] 生成与单章重跑（真实模型，约 1 分钟）")
    started = time.time()
    status, events = guest.stream(f"/api/tasks/{task_id}/generate",
                                  {"word_count": 300, "effort": "low"})
    result = next((p for n, p in events if n == "result"), {})
    check("生成成功", bool(result.get("ok")), json.dumps(result, ensure_ascii=False)[:120])
    first_file = result.get("file_id")
    check("首次生成有文件", bool(first_file), str(first_file))
    print(f"      （耗时 {time.time() - started:.1f}s）")

    nodes = [p for n, p in events if n == "node"]
    chapter_keys = sorted({p["key"] for p in nodes if str(p.get("key", "")).startswith("draft.c")})
    check("有章节节点", bool(chapter_keys), str(chapter_keys))

    if chapter_keys:
        target = chapter_keys[0]
        started = time.time()
        status, events = guest.stream(f"/api/tasks/{task_id}/nodes/{target}/retry")
        retry_result = next((p for n, p in events if n == "result"), {})
        names = [n for n, _ in events]
        check("重跑回 200", status == 200, f"got {status}")
        check("重跑成功", bool(retry_result.get("ok")), json.dumps(retry_result, ensure_ascii=False)[:120])
        check("重跑后重新出文件", bool(retry_result.get("file_id")) and retry_result["file_id"] != first_file,
              f"{first_file} → {retry_result.get('file_id')}")
        check("推了 render 事件", "render" in names, str(names))
        print(f"      （重跑耗时 {time.time() - started:.1f}s）")

        new_id = retry_result.get("file_id")
        if new_id:
            status, blob = guest.call("GET", f"/api/files/{new_id}")
            check("重跑的文件可下载", status == 200 and isinstance(blob, bytes) and blob[:2] == b"PK",
                  f"{len(blob) if isinstance(blob, bytes) else 0} 字节")

    status, events = guest.stream(f"/api/tasks/{task_id}/nodes/draft.c99/retry")
    result = next((p for n, p in events if n == "result"), {})
    check("重跑不存在章节给出理由", result.get("ok") is False and bool(result.get("error")),
          str(result.get("error"))[:40])

    status, events = guest.stream(f"/api/tasks/{task_id}/nodes/outline/retry")
    result = next((p for n, p in events if n == "result"), {})
    check("非章节节点不许单跑", result.get("ok") is False and "整篇" in str(result.get("error")),
          str(result.get("error"))[:40])

    print("\n[7] 配额只按次算")
    status, me = guest.call("GET", "/api/me")
    used = (me.get("quota") or {}).get("used")
    check("对话 1 次 + 生成 1 次 + 重跑 1 次 = 3", used == 3, f"used={used}")

    total = len(PASS) + len(FAIL)
    print(f"\n{'=' * 46}\n通过 {len(PASS)}/{total}")
    if FAIL:
        print("失败项：")
        for name in FAIL:
            print("  - " + name)
        return 1
    print("M5 材料与任务流全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
