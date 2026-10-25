"""M4 验收：Office 生成必须是真实可编辑文件。

跑法（用项目虚拟环境，里面有 python-docx / python-pptx）：
    .venv/Scripts/python.exe tests/m4_office.py

分两段：
  A 直构：调 builder 生成，再用 python-docx / python-pptx **独立**回读，
    校验字体、字号、行距、缩进、目录域、页码域、画幅比例、备注，
    并真的再存一次（可编辑性不是靠声称）。
  B 端到端：走 /generate 真调模型出一份 docx，下载回来独立回读。
"""
from __future__ import annotations

import json
import os
import sys
import time
import urllib.error
import urllib.request
import uuid
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

from app.services import capabilities, docx_builder, pptx_builder  # noqa: E402

BASE = os.environ.get("DA_BASE", "http://127.0.0.1:8020")
TMP = ROOT / "data" / "tmp" / "m4"
PASS, FAIL = [], []


def check(name: str, ok: bool, detail: str = "") -> None:
    (PASS if ok else FAIL).append(name)
    print(f"{'  OK ' if ok else '  !! '} {name}" + (f"  [{detail}]" if detail else ""))


SAMPLE = {
    "title": "行政协议诉讼的审查路径",
    "thesis": "以 2019 年司法解释为轴，重述受案范围与合法性审查的衔接",
    "sections": [
        {
            "heading": "一、问题的提出",
            "paras": [
                "行政协议进入行政诉讼受案范围之后，法院面对的第一个难题不是合法性审查本身，"
                "而是要不要先把协议拆成若干行政行为分别处理。这一处理方式直接决定了审查的顺序。",
                "既有研究多从协议性质入手，把争议归结为公法或私法的定性之争。定性当然重要，"
                "但它回答不了法官在开庭时需要先行判断的问题：这一步究竟审什么。",
            ],
        },
        {
            "heading": "二、受案范围的边界",
            "paras": [
                "司法解释把可诉的协议类型逐一列出，列举之外的协议并非当然不可诉，"
                "关键看它是否具备行政法上的权利义务内容。这一判断标准比类型清单更难操作。",
                "实践中出现的分歧集中在政府与社会资本合作项目上。这类协议的履行周期长，"
                "其间行政机关的监督行为与协议履行交织，很难按单个行为切开。",
            ],
        },
        {
            "heading": "三、审查顺序的调整",
            "paras": [
                "如果承认协议整体可诉，那么合法性审查就应当以协议约定的权利义务为对象，"
                "而不是逐条拆解行政机关的每一次意思表示。审查顺序的调整会连带影响举证责任分配。",
                "这一调整的代价是审查标准的模糊化。法院需要在尊重协议自治与监督依法行政之间"
                "找到一个可以写进判决书的尺度，这恰恰是现行规范没有提供现成答案的地方。",
            ],
        },
    ],
}


def part_a_docx() -> None:
    print("\n[A1] DOCX 直构 + 独立回读")
    out = TMP / "direct.docx"
    fmt = docx_builder.build_docx(SAMPLE, out)
    check("文件已生成", out.exists() and out.stat().st_size > 3000,
          f"{out.stat().st_size if out.exists() else 0} 字节")

    with zipfile.ZipFile(out) as zf:
        names = set(zf.namelist())
    check("是合法 OOXML 包", "word/document.xml" in names and "[Content_Types].xml" in names,
          f"{len(names)} 个部件")

    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    check("标题写入", texts and texts[0] == SAMPLE["title"], texts[0] if texts else "")
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    check("三个章节标题齐全", len(headings) == 3, str(headings))
    # 副标题也是 Normal 样式且字数够，不能按"全部 Normal"取；只取第一个标题之后的正文
    ordered = [p for p in doc.paragraphs if p.text.strip()]
    first_heading = next((i for i, p in enumerate(ordered) if p.style.name.startswith("Heading")), 0)
    body = [p for p in ordered[first_heading:]
            if p.style.name == "Normal" and docx_builder.count_words(p.text) > 20]
    check("正文段落 6 段", len(body) == 6, f"{len(body)} 段")

    first = body[0]
    run = first.runs[0]
    check("正文字体为宋体", run.font.name == fmt["body_font"], str(run.font.name))
    check("正文字号 10.5pt", round(run.font.size.pt, 1) == 10.5, f"{run.font.size.pt}pt")
    check("行距 22pt", round(first.paragraph_format.line_spacing.pt, 1) == 22.0,
          f"{first.paragraph_format.line_spacing.pt}pt")
    check("首行缩进 21pt", round(first.paragraph_format.first_line_indent.pt, 1) == 21.0,
          f"{first.paragraph_format.first_line_indent.pt}pt")

    document_xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    check("含目录域 TOC", 'TOC \\o "1-3"' in document_xml or "TOC" in document_xml)
    footer_xml = zipfile.ZipFile(out).read("word/footer1.xml").decode("utf-8") \
        if "word/footer1.xml" in zipfile.ZipFile(out).namelist() else ""
    check("含页码域 PAGE", "PAGE" in footer_xml, f"footer 部件{len(footer_xml)} 字节")

    # 可编辑性：改一个字再存一次，能被重新读出来
    doc.paragraphs[0].runs[0].text = SAMPLE["title"] + "（改）"
    again = TMP / "direct_edited.docx"
    doc.save(str(again))
    reread = Document(str(again))
    check("可编辑（另存后可回读）", reread.paragraphs[0].text.endswith("（改）"),
          reread.paragraphs[0].text[:24])

    report = docx_builder.verify_docx(out, fmt, 0, 3)
    failed = [c["name"] for c in report["checks"] if not c["ok"]]
    check("内置回读断言全过", report["ok"], f"{report['passed']}/{report['total']} 失败={failed}")

    print("\n[A2] DOCX 照样张排版")
    sample_fmt = docx_builder.read_sample_format(out)
    check("读出样张格式", bool(sample_fmt.get("body_font")),
          json.dumps({k: sample_fmt[k] for k in list(sample_fmt)[:5]}, ensure_ascii=False))
    merged = docx_builder.merge_format(sample_fmt)
    check("样张格式可合并成参数", round(merged["body_size_pt"], 1) == 10.5, str(merged["body_size_pt"]))
    out2 = TMP / "from_sample.docx"
    docx_builder.build_docx(SAMPLE, out2, sample_fmt)
    check("按样张再出一份", out2.exists() and out2.stat().st_size > 3000)


def part_a_pptx() -> None:
    print("\n[A3] PPTX 直构 + 独立回读")
    structure = {
        "title": "行政协议诉讼的审查路径",
        "slides": [
            {"layout": "cover", "title": "行政协议诉讼的审查路径",
             "subtitle": "读书会 · 2026 年 10 月", "notes": "开场三十秒说清这次要解决的问题。"},
            {"layout": "section", "title": "一、问题的提出", "notes": "过渡页，停两秒。"},
            {"layout": "bullets", "title": "审查顺序决定举证责任",
             "bullets": ["先定性再审查：回到个行为", "先审查约定：以协议为对象"],
             "notes": "强调两条路径的差别在于谁举证。"},
            {"layout": "two_col", "title": "两种路径的比较",
             "left": ["受案范围宽", "审查对象碎"], "right": ["审查对象整", "标准更模糊"],
             "notes": "左右对照，不要念稿。"},
            {"layout": "quote", "title": "一句原文",
             "subtitle": "协议的履行争议，不以行政机关的单方意思表示为限。",
             "notes": "原文照录，不解释。"},
            {"layout": "end", "title": "讨论", "notes": "留三个问题给现场。"},
        ],
    }
    out = TMP / "direct.pptx"
    info = pptx_builder.build_pptx(structure, out)
    check("文件已生成", out.exists() and out.stat().st_size > 20000,
          f"{out.stat().st_size if out.exists() else 0} 字节，模式 {info['mode']}")

    prs = Presentation(str(out))
    check("六页", len(prs.slides) == 6, f"{len(prs.slides)} 页")
    w, h = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    check("16:9 画幅", abs(w - 13.333) < 0.05 and abs(h - 7.5) < 0.05, f"{w:.2f}×{h:.2f} 英寸")

    cover_text = " ".join(s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame)
    check("封面有标题与副标题", "审查路径" in cover_text and "读书会" in cover_text, cover_text[:40])

    notes_ok = 0
    for slide in prs.slides:
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            notes_ok += 1
    check("每页都有演讲备注", notes_ok == 6, f"{notes_ok}/6")

    report = pptx_builder.verify_pptx(out, 6)
    failed = [c["name"] for c in report["checks"] if not c["ok"]]
    check("内置回读断言全过", report["ok"], f"{report['passed']}/{report['total']} 失败={failed}")

    # 版式兜底：只给标题的引文页也不能整页空白
    fallback = TMP / "fallback.pptx"
    pptx_builder.build_pptx({"title": "兜底", "slides": [
        {"layout": "cover", "title": "封面", "notes": "备注"},
        {"layout": "quote", "title": "只有标题的引文页", "notes": "备注"},
    ]}, fallback)
    fb = Presentation(str(fallback))
    fb_text = " ".join(s.text_frame.text for s in fb.slides[1].shapes if s.has_text_frame)
    check("引文页缺字段时不留空页", "只有标题" in fb_text, fb_text[:30])

    # 可编辑性
    prs.slides[1].shapes[0].text_frame.text = "改过的标题"
    again = TMP / "direct_edited.pptx"
    prs.save(str(again))
    check("可编辑（另存后可回读）",
          "改过的标题" in " ".join(s.text_frame.text for s in Presentation(str(again)).slides[1].shapes
                                  if s.has_text_frame))

    print("\n[A4] PPTX 套模板")
    tpl = TMP / "template.pptx"
    base = Presentation()
    base.slides.add_slide(base.slide_layouts[0])  # 让模板自带一页，更像真实模板
    base.save(str(tpl))
    out_tpl = TMP / "filled.pptx"
    info_tpl = pptx_builder.build_pptx(
        {"title": "套模板", "slides": [
            {"layout": "cover", "title": "模板封面", "subtitle": "副标题", "notes": "备注一"},
            {"layout": "bullets", "title": "模板正文", "bullets": ["甲", "乙"], "notes": "备注二"},
        ]},
        out_tpl, tpl,
    )
    check("走了模板模式", info_tpl.get("mode") == "template", str(info_tpl))
    filled = Presentation(str(out_tpl))
    check("在模板上追加两页", len(filled.slides) == 3, f"{len(filled.slides)} 页")
    texts = " ".join(s.text_frame.text for sl in filled.slides for s in sl.shapes if s.has_text_frame)
    check("占位符被填上内容", "模板封面" in texts and "模板正文" in texts, texts[:50])


def part_b_gate() -> None:
    print("\n[A5] 能力门槛")
    try:
        capabilities.require("docx.edit_existing")
        check("未验证能力被拦", False, "居然没报错")
    except Exception as exc:  # HTTPException
        status = getattr(exc, "status_code", 0)
        check("未验证能力返回 501", status == 501, f"{status} {getattr(exc, 'detail', '')}")
    try:
        capabilities.require("docx.create")
        check("已验证能力放行", True)
    except Exception as exc:  # noqa: BLE001
        check("已验证能力放行", False, str(exc))
    verified = {c["key"] for c in capabilities.public_matrix() if c["verified"]}
    check("矩阵里没有虚假宣称", verified == {"docx.create", "docx.read_format_sample",
                                            "pptx.create", "pptx.fill_template", "pdf.extract_text"},
          str(sorted(verified)))


class Client:
    def __init__(self):
        import http.cookiejar

        self.jar = http.cookiejar.CookieJar()
        self.opener = urllib.request.build_opener(
            urllib.request.ProxyHandler({}),  # 关掉本机 HTTP_PROXY，否则 127.0.0.1 会被代理走
            urllib.request.HTTPCookieProcessor(self.jar),
        )
        self.xff = f"10.31.0.{uuid.uuid4().int % 200 + 1}"

    def call(self, method: str, path: str, body=None):
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(BASE + path, data=data, method=method)
        req.add_header("Content-Type", "application/json")
        req.add_header("X-Forwarded-For", self.xff)
        try:
            with self.opener.open(req, timeout=300) as resp:
                raw = resp.read()
                if path.startswith("/api/files/") and "json" not in resp.headers.get("content-type", ""):
                    return resp.status, raw
                return resp.status, json.loads(raw.decode("utf-8", "replace"))
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode("utf-8", "replace")
            try:
                return exc.code, json.loads(raw)
            except json.JSONDecodeError:
                return exc.code, raw

    def stream(self, method: str, path: str, body=None) -> tuple[int, list, str]:
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(BASE + path, data=data, method=method)
        req.add_header("Content-Type", "application/json")
        req.add_header("Accept", "text/event-stream")
        req.add_header("X-Forwarded-For", self.xff)
        try:
            with self.opener.open(req, timeout=600) as resp:
                raw = resp.read().decode("utf-8", "replace")
                return resp.status, parse_sse(raw), raw
        except urllib.error.HTTPError as exc:
            return exc.code, [], exc.read().decode("utf-8", "replace")


def parse_sse(raw: str) -> list[tuple[str, dict]]:
    events = []
    for block in raw.split("\n\n"):
        if not block.strip() or block.strip().startswith(":"):
            continue
        name, payload = "message", {}
        for line in block.splitlines():
            if line.startswith("event:"):
                name = line[6:].strip()
            elif line.startswith("data:"):
                try:
                    payload = json.loads(line[5:].strip())
                except json.JSONDecodeError:
                    payload = {"raw": line[5:].strip()}
        events.append((name, payload))
    return events


def part_b_api() -> None:
    print("\n[B] 端到端：真调模型出一份 DOCX")
    admin = Client()
    status, _ = admin.call("POST", "/api/admin/login",
                           {"password": read_env().get("ADMIN_PASSWORD", "")})
    if status != 200:
        check("管理员登录", False, f"got {status}")
        return
    status, body = admin.call("POST", "/api/admin/invites",
                              {"count": 1, "max_uses": 1, "days": 7, "note": "m4"})
    code = (body.get("codes") or [""])[0]
    guest = Client()
    guest.call("POST", "/api/join", {"code": code, "nickname": "m4验收"})
    status, me = guest.call("GET", "/api/me")
    check("访客进入", status == 200, str(me.get("nickname")))

    status, task = guest.call("POST", "/api/tasks",
                              {"mode": "docs", "goal": "写一份 500 字左右的短文，"
                                                      "讲清楚行政协议诉讼里审查顺序为什么会影响举证责任"})
    task_id = (task.get("task") or {}).get("id") or task.get("id")
    check("建任务", bool(task_id), str(task_id))

    started = time.time()
    status, events, raw = guest.stream("POST", f"/api/tasks/{task_id}/generate",
                                       {"word_count": 500, "effort": "low"})
    names = [n for n, _ in events]
    check("生成流 200", status == 200, f"got {status} {raw[:120]}")
    result = next((p for n, p in events if n == "result"), {})
    check("生成成功", bool(result.get("ok")), json.dumps(result, ensure_ascii=False)[:200])
    check("有 node 进度事件", names.count("node") >= 4, f"node×{names.count('node')}")
    elapsed = time.time() - started
    print(f"      （真实生成耗时 {elapsed:.1f}s）")

    file_id = result.get("file_id")
    check("拿到了文件 id", bool(file_id), str(file_id))
    if not file_id:
        return

    status, blob = guest.call("GET", f"/api/files/{file_id}")
    check("下载 200", status == 200, f"got {status}")
    got = TMP / "api_result.docx"
    got.write_bytes(blob if isinstance(blob, bytes) else b"")
    check("下回来是 docx（PK 头）", got.read_bytes()[:2] == b"PK", f"{len(blob)} 字节")

    doc = Document(str(got))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    words = docx_builder.count_words("\n".join(paragraphs))
    check("正文有内容", words >= 200, f"{words} 字 / {len(paragraphs)} 段")
    check("有章节标题", any(p.style.name.startswith("Heading") for p in doc.paragraphs),
          str([p.text for p in doc.paragraphs if p.style.name.startswith("Heading")][:3]))
    check("有页码域", "PAGE" in zipfile.ZipFile(got).read("word/footer1.xml").decode("utf-8")
          if "word/footer1.xml" in zipfile.ZipFile(got).namelist() else False)

    report = result.get("report") or {}
    failed = [c["name"] for c in report.get("checks", []) if not c["ok"]]
    check("服务端回读断言全过", report.get("ok") is True,
          f"{report.get('passed')}/{report.get('total')} 失败={failed}")

    status, me = guest.call("GET", "/api/me")
    used = (me.get("quota") or {}).get("used")
    check("生成只扣 1 次额度", used == 1, f"used={used}（分章多次调模型）")


def read_env() -> dict:
    env = {}
    file = ROOT / ".env"
    for line in file.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            env[key.strip()] = value.strip().strip('"').strip("'")
    return env


def main() -> int:
    TMP.mkdir(parents=True, exist_ok=True)
    part_a_docx()
    part_a_pptx()
    part_b_gate()
    part_b_api()
    total = len(PASS) + len(FAIL)
    print(f"\n{'=' * 46}\n通过 {len(PASS)}/{total}")
    if FAIL:
        print("失败项：")
        for name in FAIL:
            print("  - " + name)
        return 1
    print("M4 Office 能力全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
