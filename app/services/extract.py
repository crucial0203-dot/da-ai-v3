"""材料文本抽取：PDF 逐页、DOCX 段落、纯文本直读。

PDF 提取失败或文字量极少时明确标记 has_text=False，让路由层给出
"这份 PDF 没有文字层，可能是扫描件"的具体提示，而不是静默返回空摘要。
"""
from __future__ import annotations

from pathlib import Path

from pypdf import PdfReader

MIN_USEFUL_CHARS = 40
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp"}


def extract(path: Path, name: str) -> dict:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path, name)
    if suffix == ".docx":
        return _extract_docx(path, name)
    if suffix in {".md", ".txt", ".csv"}:
        return _extract_text(path, name)
    if suffix == ".pptx":
        return _extract_pptx(path, name)
    if suffix in IMAGE_EXT:
        # 图片没有"文字层"可言，但它不是坏文件：能不能用取决于当前模型看不看图。
        # 所以单独标 image=True，让上层按模型能力决定要不要送进上下文。
        return {
            "kind": "image",
            "name": name,
            "pages": [],
            "text": "",
            "chars": 0,
            "has_text": False,
            "image": True,
            "note": "图片材料：只有能看图的模型才会读它",
        }
    return {
        "kind": "binary",
        "name": name,
        "pages": [],
        "text": "",
        "chars": 0,
        "has_text": False,
        "note": "该格式不参与文本分析，仅作为附件保存",
    }


def _extract_pdf(path: Path, name: str) -> dict:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        return _failed(name, f"PDF 解析失败：{type(exc).__name__}")
    pages: list[dict] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        if text:
            pages.append({"n": index, "text": text})
    total_chars = sum(len(p["text"]) for p in pages)
    has_text = total_chars >= MIN_USEFUL_CHARS
    return {
        "kind": "pdf",
        "name": name,
        "page_count": len(reader.pages),
        "pages": pages,
        "text": "\n\n".join(f"[第 {p['n']} 页]\n{p['text']}" for p in pages),
        "chars": total_chars,
        "has_text": has_text,
        "note": "" if has_text else "这份 PDF 没有可用文字层，可能是扫描件，需要先做 OCR",
    }


def _extract_docx(path: Path, name: str) -> dict:
    from docx import Document

    doc = Document(str(path))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paras)
    return {
        "kind": "docx",
        "name": name,
        "page_count": None,
        "pages": [{"n": 1, "text": text}] if text else [],
        "text": text,
        "chars": len(text),
        "has_text": len(text) >= MIN_USEFUL_CHARS,
        "note": "" if len(text) >= MIN_USEFUL_CHARS else "文档正文为空或过短",
    }


def _extract_pptx(path: Path, name: str) -> dict:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        return _failed(name, f"PPTX 解析失败：{type(exc).__name__}")
    chunks: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if texts:
            chunks.append(f"[第 {index} 页]\n" + "\n".join(texts))
    text = "\n\n".join(chunks)
    return {
        "kind": "pptx",
        "name": name,
        "page_count": len(prs.slides),
        "pages": [{"n": i + 1, "text": c} for i, c in enumerate(chunks)],
        "text": text,
        "chars": len(text),
        "has_text": len(text) >= MIN_USEFUL_CHARS,
        "note": "",
    }


def _extract_text(path: Path, name: str) -> dict:
    try:
        text = path.read_text(encoding="utf-8", errors="replace").strip()
    except Exception as exc:
        return _failed(name, f"读取失败：{type(exc).__name__}")
    return {
        "kind": "text",
        "name": name,
        "page_count": None,
        "pages": [{"n": 1, "text": text}] if text else [],
        "text": text,
        "chars": len(text),
        "has_text": len(text) >= MIN_USEFUL_CHARS,
        "note": "",
    }


def _failed(name: str, note: str) -> dict:
    return {
        "kind": "unknown",
        "name": name,
        "pages": [],
        "text": "",
        "chars": 0,
        "has_text": False,
        "note": note,
    }
